"""统一文件数据源读取：CSV、Excel 和 Parquet。"""
from __future__ import annotations
from collections.abc import Iterator
from pathlib import Path
import numpy as np
import pandas as pd
import pyarrow.parquet as pq

class SourceReadError(ValueError):
    pass

def _excel_chunks(source: Path, sheet: str | int | None, chunk_size: int, suffix: str) -> Iterator[pd.DataFrame]:
    """按 chunk_size 流式读取 Excel，避免一次性整表加载与逐块 copy。

    - .xlsx 走 openpyxl read_only 只读模式，按行迭代并在凑满 chunk_size 时产出 DataFrame；
    - .xls 旧格式 xlrd 不支持流式，回退整表读取后分块（保留可用性）。
    """
    if suffix == ".xls":
        try:
            frame = pd.read_excel(source, sheet_name=0 if sheet is None else sheet, dtype=str, keep_default_na=False)
        except ImportError as exc:
            raise SourceReadError("读取 .xls 需要安装可选依赖 xlrd") from exc
        except ValueError as exc:
            raise SourceReadError(f"无法读取 Excel 工作表 {sheet!r}：{exc}") from exc
        for offset in range(0, len(frame), chunk_size):
            yield frame.iloc[offset:offset + chunk_size].copy()
        return

    try:
        import openpyxl
    except ImportError as exc:
        raise SourceReadError("读取 .xlsx 需要安装可选依赖 openpyxl") from exc

    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对损坏/非 xlsx 文件抛多种异常
        raise SourceReadError(f"无法读取 Excel 工作表 {sheet!r}：{exc}") from exc

    try:
        target = 0 if sheet is None else sheet
        names = workbook.sheetnames
        if isinstance(target, int):
            if target < 0 or target >= len(names):
                raise SourceReadError(f"无法读取 Excel 工作表 {sheet!r}：工作表不存在")
            worksheet = workbook.worksheets[target]
        else:
            if target not in names:
                raise SourceReadError(f"无法读取 Excel 工作表 {sheet!r}：工作表不存在")
            worksheet = workbook[target]

        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return
        columns = ["" if value is None else str(value) for value in header]
        width = len(columns)
        buffer: list[list[str]] = []

        def _flush() -> pd.DataFrame:
            frame = pd.DataFrame(buffer, columns=columns, dtype=object)
            buffer.clear()
            return frame

        for row in rows:
            values = list(row)
            if all(value is None for value in values):
                continue  # 跳过整行空白，避免格式化残留产生的幽灵行
            if len(values) < width:
                values.extend([None] * (width - len(values)))
            buffer.append(["" if value is None else str(value) for value in values[:width]])
            if len(buffer) >= chunk_size:
                yield _flush()
        if buffer:
            yield _flush()
    finally:
        workbook.close()


def _raw_chunks(source: Path, sheet: str | int | None, chunk_size: int) -> Iterator[pd.DataFrame]:
    """按 chunk_size 分块读取原始 CSV/Excel/Parquet 文件。"""
    suffix = source.suffix.lower()
    if suffix == ".csv":
        yield from pd.read_csv(source, dtype=str, keep_default_na=False, encoding="utf-8-sig", chunksize=chunk_size, low_memory=False)
        return
    if suffix in {".xlsx", ".xls"}:
        yield from _excel_chunks(source, sheet, chunk_size, suffix)
        return
    if suffix == ".parquet":
        try:
            parquet = pq.ParquetFile(source)
            for batch in parquet.iter_batches(batch_size=chunk_size):
                yield batch.to_pandas()
        except Exception as exc:
            raise SourceReadError(f"无法读取 Parquet 文件：{exc}") from exc
        return
    raise SourceReadError(f"不支持的文件类型 {suffix or '<无后缀>'}；仅支持 .csv、.xlsx、.xls、.parquet")


def apply_derived_dimensions(chunk: pd.DataFrame, derived_dimensions: list[dict]) -> pd.DataFrame:
    """按 derived_dimensions 配置生成逻辑维度列（concat/copy），返回修改后的 chunk。

    concat 类型：按 separator 拼接 sources 中各列，空值替换为 '<空>'。
    copy 类型：将 sources[0] 列复制为新列 name。
    """
    if not derived_dimensions:
        return chunk
    for spec in derived_dimensions:
        name = spec["name"]
        dtype = spec["type"]
        sources = spec["sources"]
        missing = [s for s in sources if s not in chunk.columns]
        if missing:
            raise SourceReadError(f"派生维度 {name!r} 缺少来源列：{sorted(missing)}")
        if dtype == "copy":
            chunk[name] = chunk[sources[0]].fillna("").astype(str).str.strip()
            continue
        if dtype == "concat":
            separator = spec.get("separator", " / ")
            parts = []
            for s in sources:
                col = chunk[s].fillna("").astype(str).str.strip().replace("", "<空>")
                parts.append(col)
            result = parts[0]
            for i in range(1, len(parts)):
                result = result + separator + parts[i]
            chunk[name] = result
            continue
        raise SourceReadError(f"派生维度 {name!r} 的类型 {dtype!r} 不受支持（仅 concat/copy）")
    return chunk


def read_scoped_file(
    source: Path,
    data_source: dict,
    numeric: set[str],
    start,
    end,
    chunk_size: int,
    date_logical: str,
    derived_dimensions: list[dict] | None = None,
):
    """读取数据源并返回 (原始行数, 过滤并映射后的 chunk) 迭代器。

    - date_logical：日期逻辑列名（替代硬编码 opration_date）。
    - derived_dimensions：派生逻辑维配置，在映射后、时间过滤前应用。
    """
    mapping = data_source.get("field_mapping", {})
    rename = {physical: logical for logical, physical in mapping.items()}
    required_physical = set(mapping.values())
    for raw in _raw_chunks(source, data_source.get("sheet"), chunk_size):
        raw_count = len(raw)
        raw.columns = [str(column).strip() for column in raw.columns]
        missing = required_physical - set(raw.columns)
        if missing:
            raise SourceReadError(f"数据源缺少映射字段：{sorted(missing)}")
        chunk = raw[list(required_physical)].rename(columns=rename)
        for column in chunk.columns:
            chunk[column] = chunk[column].fillna("").astype(str).str.strip()
        for column in numeric & set(chunk.columns):
            chunk[column] = pd.to_numeric(chunk[column].str.replace(",", "", regex=False).str.replace("￥", "", regex=False).replace("", np.nan), errors="coerce")
        if date_logical not in chunk.columns:
            raise SourceReadError(f"日期逻辑列 {date_logical!r} 不存在于映射结果中")
        chunk = apply_derived_dimensions(chunk, derived_dimensions or [])
        dates = pd.to_datetime(chunk[date_logical].replace("", pd.NaT), errors="coerce")
        chunk["analysis_month"] = dates.dt.strftime("%Y-%m").fillna("")
        chunk = chunk.loc[(dates >= start) & (dates < end)].copy()
        if not chunk.empty:
            chunk["analysis_date"] = dates.loc[chunk.index]
        yield raw_count, chunk
